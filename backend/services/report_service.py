import logging
import io
from decimal import Decimal
from typing import List, Dict, Any
from datetime import date

from sqlalchemy.orm import Session
from sqlalchemy import extract
from fastapi import HTTPException
from fastapi.responses import StreamingResponse

from models.models import Transaction, GSTReturn, TransactionType
from services.gst_service import compute_gstr1_summary, compute_gstr3b_summary

logger = logging.getLogger("autogst.reports")


def get_transactions_for_period(
    db: Session, user_id: int, month: int, year: int
) -> List[Transaction]:
    """Fetch transactions for a given month/year."""
    return (
        db.query(Transaction)
        .filter(
            Transaction.user_id == user_id,
            extract("month", Transaction.transaction_date) == month,
            extract("year", Transaction.transaction_date) == year,
        )
        .all()
    )


def generate_report_data(
    db: Session, user_id: int, report_type: str, month: int, year: int
) -> Dict[str, Any]:
    """Generate GST report data for the given period."""
    transactions = get_transactions_for_period(db, user_id, month, year)

    if report_type == "GSTR-1":
        data = compute_gstr1_summary(transactions)
    elif report_type == "GSTR-3B":
        data = compute_gstr3b_summary(transactions)
    else:
        raise HTTPException(400, f"Unknown report type: {report_type}")

    data["metadata"] = {
        "report_type": report_type,
        "period": f"{year}-{month:02d}",
        "transaction_count": len(transactions),
        "generated_at": str(date.today()),
    }
    return data


def generate_pdf_report(data: Dict[str, Any], report_type: str) -> bytes:
    """Generate a PDF report using ReportLab."""
    try:
        from reportlab.lib.pagesizes import A4
        from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
        from reportlab.lib.units import cm
        from reportlab.lib import colors
        from reportlab.platypus import (
            SimpleDocTemplate, Table, TableStyle, Paragraph,
            Spacer, HRFlowable
        )

        buffer = io.BytesIO()
        doc = SimpleDocTemplate(
            buffer, pagesize=A4,
            topMargin=2*cm, bottomMargin=2*cm,
            leftMargin=2*cm, rightMargin=2*cm
        )
        styles = getSampleStyleSheet()
        story = []

        # Title
        title_style = ParagraphStyle(
            "title", parent=styles["Title"],
            fontSize=18, textColor=colors.HexColor("#1a1a2e"),
            spaceAfter=6,
        )
        story.append(Paragraph(f"AutoGST Pro — {report_type} Report", title_style))

        meta = data.get("metadata", {})
        story.append(Paragraph(f"Period: {meta.get('period', 'N/A')} | Generated: {meta.get('generated_at', 'N/A')}", styles["Normal"]))
        story.append(Spacer(1, 0.5*cm))
        story.append(HRFlowable(width="100%", thickness=1, color=colors.HexColor("#e0e0e0")))
        story.append(Spacer(1, 0.5*cm))

        # Summary table
        if report_type == "GSTR-1":
            summary = data.get("summary", {})
            summary_rows = [
                ["Metric", "Amount (₹)"],
                ["Total Taxable Value", f"{summary.get('total_taxable_value', 0):,.2f}"],
                ["Total CGST", f"{summary.get('total_cgst', 0):,.2f}"],
                ["Total SGST", f"{summary.get('total_sgst', 0):,.2f}"],
                ["Total IGST", f"{summary.get('total_igst', 0):,.2f}"],
                ["Total Tax", f"{summary.get('total_tax', 0):,.2f}"],
            ]
        else:  # GSTR-3B
            payment = data.get("6_payment_of_tax", {})
            summary_rows = [
                ["Metric", "Amount (₹)"],
                ["Total Output Tax", f"{payment.get('total_output_tax', 0):,.2f}"],
                ["Total ITC", f"{payment.get('total_itc', 0):,.2f}"],
                ["Net Tax Payable", f"{payment.get('net_payable', 0):,.2f}"],
            ]

        table = Table(summary_rows, colWidths=[10*cm, 7*cm])
        table.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1a1a2e")),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
            ("FONTSIZE", (0, 0), (-1, -1), 11),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f5f5f5")]),
            ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#cccccc")),
            ("ALIGN", (1, 0), (-1, -1), "RIGHT"),
            ("LEFTPADDING", (0, 0), (-1, -1), 8),
            ("RIGHTPADDING", (0, 0), (-1, -1), 8),
            ("TOPPADDING", (0, 0), (-1, -1), 6),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
        ]))
        story.append(table)

        doc.build(story)
        return buffer.getvalue()
    except Exception as e:
        logger.error("PDF generation failed: %s", str(e))
        raise HTTPException(500, f"PDF generation failed: {str(e)}")


def generate_excel_report(
    db: Session, user_id: int, month: int, year: int, report_type: str
) -> bytes:
    """Generate Excel report using openpyxl."""
    try:
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
        from openpyxl.utils import get_column_letter

        transactions = get_transactions_for_period(db, user_id, month, year)

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = f"{report_type} {year}-{month:02d}"

        # Header style
        header_fill = PatternFill("solid", fgColor="1a1a2e")
        header_font = Font(color="FFFFFF", bold=True, size=11)
        border = Border(
            left=Side(style="thin"), right=Side(style="thin"),
            top=Side(style="thin"), bottom=Side(style="thin")
        )

        headers = [
            "Date", "Description", "Type", "Party", "GSTIN",
            "Invoice No.", "HSN/SAC", "Taxable Amount",
            "CGST", "SGST", "IGST", "Total GST", "Total Amount",
            "Category", "Anomaly"
        ]

        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=header)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center")
            cell.border = border

        # Data rows
        alt_fill = PatternFill("solid", fgColor="F5F5F5")
        for row_idx, txn in enumerate(transactions, 2):
            fill = alt_fill if row_idx % 2 == 0 else PatternFill()
            values = [
                str(txn.transaction_date),
                txn.description,
                txn.transaction_type.value,
                txn.party_name or "",
                txn.party_gstin or "",
                txn.invoice_number or "",
                txn.hsn_sac_code or "",
                float(txn.taxable_amount or 0),
                float(txn.cgst_amount or 0),
                float(txn.sgst_amount or 0),
                float(txn.igst_amount or 0),
                float(txn.total_gst or 0),
                float(txn.total_amount or 0),
                txn.ml_category or "",
                "Yes" if txn.is_anomaly else "No",
            ]
            for col, value in enumerate(values, 1):
                cell = ws.cell(row=row_idx, column=col, value=value)
                cell.fill = fill
                cell.border = border
                if isinstance(value, float):
                    cell.number_format = "#,##0.00"

        # Auto-fit columns
        for col in ws.columns:
            max_len = max((len(str(cell.value)) for cell in col if cell.value), default=10)
            ws.column_dimensions[get_column_letter(col[0].column)].width = min(max_len + 4, 30)

        # Summary sheet
        ws2 = wb.create_sheet("Summary")
        summary_data = (
            compute_gstr1_summary(transactions)
            if report_type == "GSTR-1"
            else compute_gstr3b_summary(transactions)
        )

        ws2["A1"] = f"{report_type} Summary — {year}-{month:02d}"
        ws2["A1"].font = Font(bold=True, size=14)

        # Write summary key-value pairs
        summary_flat = {}
        if report_type == "GSTR-1":
            summary_flat = summary_data.get("summary", {})
        else:
            for section, vals in summary_data.items():
                if isinstance(vals, dict):
                    for k, v in vals.items():
                        summary_flat[f"{section} — {k}"] = v

        for row_i, (k, v) in enumerate(summary_flat.items(), start=3):
            ws2.cell(row=row_i, column=1, value=k)
            ws2.cell(row=row_i, column=2, value=v)

        buffer = io.BytesIO()
        wb.save(buffer)
        return buffer.getvalue()
    except Exception as e:
        logger.error("Excel generation failed: %s", str(e))
        raise HTTPException(500, f"Excel generation failed: {str(e)}")
